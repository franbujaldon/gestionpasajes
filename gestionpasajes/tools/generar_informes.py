from pathlib import Path
from copy import deepcopy

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE_PERSONAL = ROOT / "Informe personal" / "Informe personal MongoDB gestion de pasajes.docx"
OUTPUT_PERSONAL = ROOT / "Informe personal" / "Informe personal MongoDB gestion de pasajes - version 2.docx"
OUTPUT_PROFESOR = ROOT / "Informe para entrega" / "Informe tecnico Proyecto MongoDB gestion de pasajes.docx"
ASSETS = ROOT / "gestionpasajes" / "informes" / "assets"

AZUL = "17365D"
AZUL_MEDIO = "2F5597"
AZUL_CLARO = "D9EAF7"
GRIS = "F2F2F2"
VERDE = "E2F0D9"
BLANCO = "FFFFFF"
NEGRO = RGBColor(31, 31, 31)


def asegurar_carpetas():
    OUTPUT_PROFESOR.parent.mkdir(parents=True, exist_ok=True)
    ASSETS.mkdir(parents=True, exist_ok=True)


def configurar_fuentes(doc):
    # El informe original guarda los nombres internos Heading con mayúscula.
    # python-docx espera los nombres incorporados en minúscula para add_heading().
    for style in doc.styles:
        if style.style_id in {"Heading1", "Heading2", "Heading3", "Heading4", "Heading5", "Heading6", "Heading7", "Heading8", "Heading9"}:
            name_node = style._element.find(qn("w:name"))
            if name_node is not None:
                name_node.set(qn("w:val"), name_node.get(qn("w:val")).lower())
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = NEGRO
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for nombre, tamano, color in [
        ("Title", 24, AZUL),
        ("Subtitle", 13, AZUL_MEDIO),
        ("Heading 1", 16, AZUL),
        ("Heading 2", 12.5, AZUL_MEDIO),
        ("Heading 3", 10.5, AZUL_MEDIO),
    ]:
        estilo = doc.styles[nombre]
        estilo.font.name = "Aptos Display" if nombre != "Normal" else "Aptos"
        estilo.font.size = Pt(tamano)
        estilo.font.color.rgb = RGBColor.from_string(color)
        estilo.font.bold = nombre != "Subtitle"
        estilo.paragraph_format.keep_with_next = True
        estilo.paragraph_format.space_before = Pt(10 if nombre != "Title" else 0)
        estilo.paragraph_format.space_after = Pt(5)
    if "Código" not in [s.name for s in doc.styles]:
        estilo = doc.styles.add_style("Código", WD_STYLE_TYPE.PARAGRAPH)
        estilo.font.name = "Consolas"
        estilo.font.size = Pt(8)
        estilo.paragraph_format.space_before = Pt(3)
        estilo.paragraph_format.space_after = Pt(5)
        estilo.paragraph_format.keep_together = True


def configurar_pagina(doc):
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.7)
        section.left_margin = Cm(1.9)
        section.right_margin = Cm(1.9)
        section.header_distance = Cm(0.7)
        section.footer_distance = Cm(0.7)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=75, start=90, bottom=75, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def no_break_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_table(table, header=True, zebra=True, font_size=8.5):
    table.style = "Table Grid"
    table.autofit = True
    for ri, row in enumerate(table.rows):
        no_break_row(row)
        if ri == 0 and header:
            repeat_table_header(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ri == 0 and header:
                set_cell_shading(cell, AZUL)
            elif zebra and ri % 2 == 0:
                set_cell_shading(cell, "EAF1F8")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
                    if ri == 0 and header:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, text in enumerate(headers):
        table.rows[0].cells[i].text = str(text)
    for row_data in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row_data):
            cells[i].text = str(text)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    format_table(table, font_size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(2)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(2)


def add_code(doc, code):
    p = doc.add_paragraph(style="Código")
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.right_indent = Cm(0.35)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(code.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), GRIS)
    p_pr.append(shd)
    return p


def add_note(doc, title, text, color=AZUL_CLARO):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    p = cell.paragraphs[0]
    r = p.add_run(f"{title}: ")
    r.bold = True
    p.add_run(text)
    format_table(table, header=False, zebra=False, font_size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = False
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(89, 89, 89)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Página ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def configurar_encabezado_pie(doc, etiqueta):
    for section in doc.sections:
        hp = section.header.paragraphs[0]
        hp.text = etiqueta
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in hp.runs:
            r.font.name = "Aptos"
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(100, 100, 100)
        fp = section.footer.paragraphs[0]
        fp.clear()
        add_page_field(fp)
        for r in fp.runs:
            r.font.name = "Aptos"
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(100, 100, 100)


def crear_diagramas():
    regular_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    bold_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
    font = ImageFont.truetype(regular_path, 25)
    small = ImageFont.truetype(regular_path, 20)
    tiny = ImageFont.truetype(regular_path, 16)
    bold = ImageFont.truetype(bold_path, 25)
    title = ImageFont.truetype(bold_path, 38)

    def center_text(draw, box, text, used_font, fill="#17365D", spacing=5):
        left, top, right, bottom = box
        bbox = draw.multiline_textbbox((0, 0), text, font=used_font, align="center", spacing=spacing)
        width = bbox[2] - bbox[0]
        height = bbox[3] - bbox[1]
        draw.multiline_text(((left + right - width) / 2, (top + bottom - height) / 2 - bbox[1]), text, font=used_font, fill=fill, align="center", spacing=spacing)

    def arrow(draw, start, end, color="#2F5597", width=5):
        draw.line([start, end], fill=color, width=width)
        import math
        angle = math.atan2(end[1] - start[1], end[0] - start[0])
        size = 18
        a1 = angle + math.pi * 0.82
        a2 = angle - math.pi * 0.82
        p1 = (end[0] + size * math.cos(a1), end[1] + size * math.sin(a1))
        p2 = (end[0] + size * math.cos(a2), end[1] + size * math.sin(a2))
        draw.polygon([end, p1, p2], fill=color)

    # Arquitectura física de la demostración.
    img = Image.new("RGB", (2200, 720), "white")
    draw = ImageDraw.Draw(img)
    center_text(draw, (0, 35, 2200, 115), "Arquitectura local de demostración", title)
    cajas = [
        (80, 260, 460, 520, "Navegador\nInterfaz web", "#D9EAF7"),
        (610, 260, 1010, 520, "Servidor local\nPython / HTTP", "#E2F0D9"),
        (1160, 260, 1540, 520, "mongosh\nmongo_api.js", "#FFF2CC"),
        (1690, 260, 2150, 520, "MongoDB 8.0.32\nWiredTiger", "#FCE4D6"),
    ]
    for x1, y1, x2, y2, text, color in cajas:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=22, fill=color, outline="#17365D", width=5)
        center_text(draw, (x1, y1, x2, y2), text, bold)
    for a, b, label in [((460, 390), (610, 390), "HTTP/JSON"), ((1010, 390), (1160, 390), "comandos"), ((1540, 390), (1690, 390), "BSON")]:
        arrow(draw, a, b)
        center_text(draw, (a[0], 315, b[0], 365), label, tiny, fill="#595959")
    img.save(ASSETS / "arquitectura.png")

    # Modelo documental resumido con la colección temporal.
    img = Image.new("RGB", (2400, 1500), "white")
    draw = ImageDraw.Draw(img)
    center_text(draw, (0, 30, 2400, 120), "Modelo documental implementado", title)
    center_text(draw, (0, 115, 2400, 180), "Las flechas son referencias ObjectId; cada caja resume datos embebidos.", small, fill="#595959")
    nodes = {
        "terminales": (80, 245, 470, 430, "terminales\nGeoJSON · 2dsphere"),
        "recorridos": (630, 245, 1020, 430, "recorridos\nparadas + tramos"),
        "colectivos": (1180, 245, 1570, 430, "colectivos\nasientos físicos"),
        "usuarios": (1930, 245, 2320, 430, "usuarios\ncuenta y rol"),
        "viajes": (900, 620, 1320, 825, "viajes\nsnapshots + inventario"),
        "pasajeros": (1880, 620, 2320, 825, "pasajeros\npersona que viaja"),
        "reservas": (300, 1030, 720, 1235, "reservas\nbloqueo + TTL"),
        "pasajes": (1570, 1030, 2020, 1235, "pasajes\nventa + pago + snapshots"),
        "seguimiento": (900, 1260, 1320, 1445, "seguimientoviajes\ntime series"),
    }
    for key, (x1, y1, x2, y2, label) in nodes.items():
        color = "#E2F0D9" if key in {"reservas", "pasajes", "seguimiento"} else "#D9EAF7"
        if key == "viajes":
            color = "#FFF2CC"
        draw.rounded_rectangle((x1, y1, x2, y2), radius=20, fill=color, outline="#17365D", width=5)
        center_text(draw, (x1, y1, x2, y2), label, bold)

    def center(k):
        x1, y1, x2, y2, _ = nodes[k]
        return ((x1 + x2) // 2, (y1 + y2) // 2)

    relaciones = [
        ("terminales", "recorridos", "terminalId", (-5, -45)),
        ("recorridos", "viajes", "recorridoId", (-100, 0)),
        ("colectivos", "viajes", "colectivoId", (100, 0)),
        ("usuarios", "pasajeros", "pasajeroId opcional", (110, 0)),
        ("pasajeros", "reservas", "pasajeroId", (-50, 60)),
        ("viajes", "reservas", "viajeId", (-50, 0)),
        ("viajes", "pasajes", "viajeId", (30, 0)),
        ("pasajeros", "pasajes", "pasajeroId", (45, 0)),
        ("viajes", "seguimiento", "meta.viajeId", (110, 0)),
    ]
    for a, b, label, offset in relaciones:
        start, end = center(a), center(b)
        arrow(draw, start, end, color="#6B7A8F", width=4)
        mx, my = (start[0] + end[0]) // 2 + offset[0], (start[1] + end[1]) // 2 + offset[1]
        bbox = draw.textbbox((0, 0), label, font=tiny)
        pad = 6
        draw.rectangle((mx - (bbox[2] - bbox[0]) / 2 - pad, my - 16 - pad, mx + (bbox[2] - bbox[0]) / 2 + pad, my + 16 + pad), fill="white")
        draw.text((mx - (bbox[2] - bbox[0]) / 2, my - 16), label, font=tiny, fill="#595959")
    img.save(ASSETS / "modelo_documental_actualizado.png")


def cambiar_parrafo(doc, texto_inicial, reemplazo):
    for p in doc.paragraphs:
        if p.text.strip().startswith(texto_inicial):
            p.text = reemplazo
            return True
    return False


def copiar_formato_fila(origen, destino):
    for i, cell in enumerate(origen.cells):
        destino.cells[i]._tc.get_or_add_tcPr().append(deepcopy(cell._tc.get_or_add_tcPr()))


def actualizar_informe_personal():
    doc = Document(SOURCE_PERSONAL)
    configurar_fuentes(doc)
    configurar_pagina(doc)
    configurar_encabezado_pie(doc, "Informe personal MongoDB · gestión de pasajes")

    reemplazos = [
        ("Versión 1", "Versión 2 — 23 de octubre de 2026"),
        ("El contenido técnico de las colecciones representa un modelo preliminar", "El contenido técnico documenta el modelo implementado y verificado. Los ejemplos se corresponden con la base gestionpasajes, los scripts ejecutables, la interfaz local y las pruebas de integridad, rendimiento y restauración realizadas."),
        ("5.  Diseño preliminar", "5.  Diseño implementado de la base gestionpasajes"),
        ("6.  Ejemplos de documentos para las ocho colecciones", "6.  Ejemplos de documentos para las nueve colecciones"),
        ("Seguimiento GPS en tiempo real.", "Integración con dispositivos GPS reales y transmisión continua en producción. Se incluye, con fines académicos, una serie temporal de telemetría simulada."),
        ("4 1 Colecciones preliminares", "4 1 Colecciones implementadas"),
        ("Figura 2  Modelo documental preliminar", "Figura 2  Modelo documental de la base gestionpasajes"),
        ("12 2 Política preliminar", "12 2 Política aplicada y recomendada"),
        ("18 Próxima revisión", "18 Cierre de la versión preliminar"),
        ("La próxima versión deberá cerrar el modelo físico", "El modelo físico, los validadores, la carga, las operaciones atómicas, los índices y las pruebas se implementaron. Las secciones siguientes registran los resultados finales y amplían la relación entre la teoría de la cátedra y las decisiones del proyecto."),
    ]
    for inicial, nuevo in reemplazos:
        cambiar_parrafo(doc, inicial, nuevo)

    # Estado real del proyecto.
    t = doc.tables[0]
    actualizaciones = {
        "Colecciones": ("Implementado", "Nueve colecciones; una de ellas es time series"),
        "Modelo físico": ("Implementado", "Validadores, referencias, embebidos, snapshots e índices"),
        "Interfaz": ("Implementado", "Demostración local conectada a MongoDB"),
    }
    for row in t.rows[1:]:
        key = row.cells[0].text.strip()
        if key in actualizaciones:
            row.cells[1].text, row.cells[2].text = actualizaciones[key]
    format_table(t)

    t = doc.tables[1]
    estados = {
        "Necesidad y solución": "Completado",
        "Modelado": "Completado",
        "Implementación": "Completado y probado",
        "Consultas avanzadas": "Diez consultas ejecutables",
        "Rendimiento": "Verificado con explain",
        "Resguardos": "Dump y restauración verificados",
        "Exposición": "Guion técnico preparado",
    }
    for row in t.rows[1:]:
        if row.cells[0].text.strip() in estados:
            row.cells[2].text = estados[row.cells[0].text.strip()]
    format_table(t)

    t = doc.tables[4]
    nueva = t.add_row()
    nueva.cells[0].text = "seguimientoviajes"
    nueva.cells[1].text = "Telemetría histórica por fecha"
    nueva.cells[2].text = "meta de viaje/colectivo y mediciones"
    format_table(t)

    t = doc.tables[7]
    extras = [
        ("viajes", "{ paresVenta: 1, estado: 1, fechaSalida: 1 }", "Búsqueda comercial por origen, destino y fecha"),
        ("pasajes", "{ pasajeroSnapshot.tipoDocumento: 1, numeroDocumento: 1 }", "Búsqueda directa de pasajes por DNI"),
        ("seguimientoviajes", "{ meta.viajeId: 1, fecha: 1 }", "Rango temporal de un viaje"),
    ]
    for data in extras:
        row = t.add_row()
        for i, value in enumerate(data):
            row.cells[i].text = value
    format_table(t, font_size=7.8)

    doc.add_page_break()
    doc.add_heading("19 Actualización final de la implementación", level=1)
    doc.add_paragraph(
        "Esta actualización convierte el documento de estudio inicial en una memoria de aprendizaje completa. "
        "Distingue tres planos: la teoría de las clases, la decisión de diseño y la evidencia obtenida al ejecutar "
        "MongoDB. Esa separación es útil para defender el proyecto: una característica no se justifica porque MongoDB "
        "la permita, sino porque responde a una consulta, una regla o un requisito concreto."
    )
    add_table(doc, ["Elemento", "Resultado implementado", "Evidencia"], [
        ("Motor", "MongoDB Community 8.0.32 / WiredTiger", "serverStatus y verificación final"),
        ("Base", "gestionpasajes", "una base, nueve colecciones lógicas"),
        ("Datos", "124 documentos lógicos", "incluye 24 mediciones temporales"),
        ("Esquema", "JSON Schema estricto en ocho colecciones", "validationAction: error"),
        ("Índices", "35 entradas en total", "8 _id, 1 automático time series y 26 definidos"),
        ("Consultas", "diez demostraciones avanzadas", "find, arrays, aggregate, lookup, GIS, time series y explain"),
        ("Respaldo", "mongodump restaurado con éxito", "conteos, tipos e índices equivalentes"),
        ("Interfaz", "aplicación local funcional", "búsqueda, disponibilidad, reservas e informes"),
    ], widths=[3.2, 6.7, 6.6])

    doc.add_heading("19 1 Qué colección es la principal", level=2)
    doc.add_paragraph(
        "MongoDB no posee una propiedad formal llamada colección principal. En este dominio, viajes es el agregado "
        "operativo central porque reúne fecha, recorrido, colectivo, inventario y ocupaciones que la venta necesita leer "
        "juntos. Pasajes es el centro transaccional e histórico: cada documento prueba una venta confirmada. Esta respuesta "
        "es más precisa que elegir una sola colección sin explicar el punto de vista. La centralidad se demuestra por el "
        "flujo de consultas y por la cantidad de referencias hacia viajes, no mediante una opción especial de MongoDB."
    )
    add_note(doc, "Respuesta breve para la exposición", "Viajes es la colección principal para operar la venta; pasajes es la principal para auditar y reportar las ventas. MongoDB no obliga a declarar una colección raíz.")

    doc.add_page_break()
    doc.add_heading("20 Teoría integrada de las presentaciones", level=1)
    doc.add_paragraph(
        "Las presentaciones de la cátedra son guías temáticas. A continuación se registra qué concepto aporta cada una "
        "y cómo se materializa en gestionpasajes. Los temas históricos o de producto que no afectan esta solución se "
        "mantienen como contexto, mientras que los conceptos aplicables se enlazan con evidencia concreta."
    )
    add_table(doc, ["Presentación", "Teoría relevante", "Aplicación en el proyecto"], [
        ("0 SQL Server", "modelo relacional, transacciones, seguridad, optimizador, integridad y backup", "sirve como contraste: referencias no impuestas, documentos embebidos, explain y resguardo BSON"),
        ("1 Infraestructura", "stack, capas, virtualización, cloud, clustering, Big Data y BI/ETL", "arquitectura navegador–servidor–mongosh–MongoDB; posible migración a Atlas"),
        ("2 GIS", "dato espacial, coordenadas, proyecciones, vector frente a raster", "terminales como puntos GeoJSON y búsquedas con 2dsphere"),
        ("3 NoSQL", "familias NoSQL, documentos BSON/JSON y escalabilidad horizontal", "motor documental y modelo orientado a agregados"),
        ("4 Introducción MongoDB", "documentos, flexibilidad, replicación y grandes volúmenes", "estructura BSON flexible con validación selectiva"),
        ("5 Instalación", "mongod, mongosh, archivos, MVCC, checkpoints, compresión y journal", "servidor local 8.0.32 con WiredTiger y scripts mongosh"),
        ("6 Modelado", "el diseño de la aplicación y la navegabilidad guían el modelo denormalizado", "paresVenta, snapshots e inventario nacen de consultas de interfaz"),
        ("7 Operación", "CRUD, tipos, arrays, relaciones, aggregate, lookup y time series", "operaciones ejecutables, pipelines y seguimientoviajes"),
        ("8 Índices", "índices simples, compuestos, multikey, geoespaciales, hashed y explain", "índices según consultas; medición IXSCAN frente a COLLSCAN"),
        ("9 Backup", "mongoexport/import frente a mongodump/mongorestore", "respaldo BSON completo y prueba de restauración"),
        ("9b Cloud", "Atlas, clústeres, réplicas, UI/CLI, seguridad y administración", "alternativa de despliegue; la entrega actual permanece local"),
    ], widths=[3.2, 6.5, 7.2], font_size=7.7)

    doc.add_heading("20 1 Relacional y documental: la diferencia importante", level=2)
    doc.add_paragraph(
        "Una base relacional distribuye hechos entre tablas normalizadas y reconstruye una vista mediante JOIN. MongoDB "
        "almacena documentos BSON que pueden contener objetos y arreglos. NoSQL no significa ausencia de relaciones: "
        "significa que la relación puede resolverse embebiendo datos, guardando ObjectId o aplicando $lookup. La decisión "
        "se toma según cardinalidad, patrón de acceso, crecimiento y necesidad de actualización independiente."
    )
    add_table(doc, ["Pregunta", "Si la respuesta es sí", "Decisión habitual"], [
        ("¿Los datos se consultan siempre juntos?", "evitar lectura adicional", "embeber"),
        ("¿El hijo tiene tamaño acotado?", "el documento no crecerá sin límite", "embeber"),
        ("¿Tiene vida propia o se comparte?", "varios documentos lo reutilizan", "referenciar"),
        ("¿Debe conservarse el valor histórico?", "el dato vigente puede cambiar", "snapshot denormalizado"),
        ("¿La relación se consulta ocasionalmente?", "no conviene duplicar todo", "$lookup o consultas separadas"),
    ], widths=[6.0, 5.2, 5.4])

    doc.add_heading("20 2 ACID, atomicidad y consistencia en este diseño", level=2)
    doc.add_paragraph(
        "La atomicidad documental es una fortaleza clave: una actualización que modifica un viaje completo se aplica o no "
        "se aplica. reservarAsiento usa un filtro condicional que comprueba superposición y un $push atómico sobre el mismo "
        "documento. Así, dos pedidos concurrentes no pueden ganar el mismo asiento y tramo. Insertar luego la reserva es una "
        "segunda operación; si falla, el código aplica una compensación con $pull. Para garantías multidocumento estrictas "
        "podría emplearse una transacción en un replica set, pero la solución académica prioriza el agregado viaje."
    )

    doc.add_heading("20 3 BSON y elección de tipos", level=2)
    add_bullets(doc, [
        "ObjectId identifica documentos y referencias sin depender de un entero centralizado.",
        "Date conserva instantes; los filtros por rango usan $gte y $lt sobre valores temporales reales.",
        "Decimal128 representa importes sin errores binarios propios de double.",
        "Array modela paradas, tramos, asientos, ocupaciones y pares de venta.",
        "Documento embebido agrupa contacto, domicilio, pago, embarque y snapshots.",
        "GeoJSON Point almacena coordenadas en orden longitud, latitud y habilita 2dsphere.",
    ])

    doc.add_heading("21 Modelo documental físico", level=1)
    doc.add_picture(str(ASSETS / "modelo_documental_actualizado.png"), width=Inches(6.8))
    add_caption(doc, "Figura 3. Modelo documental final: referencias entre colecciones y datos embebidos.")
    add_table(doc, ["Colección", "Patrón", "Justificación"], [
        ("usuarios", "referencia opcional a pasajero", "cuenta y persona no son lo mismo"),
        ("pasajeros", "contacto y domicilio embebidos", "datos acotados de la persona"),
        ("terminales", "GeoJSON embebido", "punto espacial propio de la terminal"),
        ("recorridos", "paradas y tramos embebidos", "secuencia ordenada consultada como unidad"),
        ("colectivos", "asientos embebidos", "configuración física acotada"),
        ("viajes", "snapshots e inventario embebidos", "lectura rápida y control atómico"),
        ("reservas", "documento independiente con TTL", "ciclo de vida temporal y consulta por pasajero"),
        ("pasajes", "snapshots de pasajero, viaje y pago", "historia de una venta aun si cambian maestros"),
        ("seguimientoviajes", "time series con meta estable", "muchas mediciones ordenadas por fecha"),
    ], widths=[3.2, 5.3, 8.2], font_size=8)

    doc.add_heading("21 1 Redundancia controlada", level=2)
    doc.add_paragraph(
        "La redundancia no es accidental. paresVenta evita calcular todas las combinaciones de paradas en cada búsqueda; "
        "pasajeroSnapshot permite hallar pasajes por DNI; viajeSnapshot conserva código, salida, recorrido y colectivo; "
        "recorridoSnapshot evita que una modificación futura reescriba el pasado. Los snapshots históricos no se sincronizan: "
        "representan lo que era cierto al vender. Los campos derivados vigentes sí deben regenerarse cuando cambia su fuente."
    )

    doc.add_heading("22 Operaciones y consultas avanzadas", level=1)
    doc.add_paragraph(
        "El catálogo scripts/11_consultas_presentacion.js reúne diez consultas de solo lectura. My Queries es una función "
        "local de Compass: guarda filtros y pipelines en el entorno del usuario, pero no crea una colección en la base ni "
        "viaja dentro de mongodump. Por eso existe también una copia portable en el repositorio."
    )
    add_table(doc, ["N.º", "Consulta", "Conceptos"], [
        ("1", "viajes Mendoza–Córdoba", "find, multikey, rango, sort, índice compuesto"),
        ("2", "asientos libres por tramo", "$filter, $map, $size, arreglos embebidos"),
        ("3", "pasajes por DNI", "snapshot, proyección e índice"),
        ("4", "manifiesto", "$lookup, $unwind, $project y $sort"),
        ("5", "recaudación", "$match, $group, $sum, $avg y $round"),
        ("6", "ocupación por viaje", "cálculo sobre inventario embebido"),
        ("7", "terminales cercanas", "$geoNear y 2dsphere"),
        ("8", "reservas y TTL", "agrupación por estado e índice parcial"),
        ("9", "velocidad por 15 minutos", "time series, $dateTrunc y $avg"),
        ("10", "plan de ejecución", "explain('executionStats')"),
    ], widths=[1.0, 6.3, 9.4])
    add_code(doc, '''db.pasajes.aggregate([
  { $match: { "pago.estado": "aprobado" } },
  { $group: {
      _id: { canal: "$venta.canal", medio: "$pago.medio" },
      cantidad: { $sum: 1 },
      recaudacion: { $sum: "$pago.importe" }
  }},
  { $sort: { recaudacion: -1 } }
])''')
    doc.add_paragraph(
        "Lectura: $match reduce el conjunto; $group define una clave compuesta y acumula; $sum con valor 1 cuenta; "
        "$sum sobre el importe recauda; $sort ordena el resultado. Un pipeline se interpreta de izquierda a derecha: "
        "la salida de una etapa es la entrada de la siguiente."
    )

    doc.add_heading("23 Índices y rendimiento", level=1)
    doc.add_paragraph(
        "Un índice mantiene una estructura ordenada que reduce documentos examinados, pero ocupa memoria y disco y agrega "
        "trabajo a inserciones, actualizaciones y eliminaciones. Por eso cada índice del proyecto responde a una consulta, "
        "una unicidad, un vencimiento o una operación geoespacial. Los índices sobre arreglos se vuelven multikey de forma "
        "automática; los compuestos deben respetar el prefijo útil de igualdad, orden y rango."
    )
    add_table(doc, ["Prueba", "Devueltos", "Claves examinadas", "Documentos examinados", "Etapa"], [
        ("Con ix_viajes_parVenta_estado_fecha", "3", "3", "3", "IXSCAN"),
        ("Forzada con {$natural: 1}", "3", "0", "12", "COLLSCAN"),
    ], widths=[7.0, 2.1, 2.7, 3.1, 2.0])
    add_note(doc, "Interpretación correcta", "La colección de prueba es pequeña, de modo que executionTimeMillis puede ser 0 o variar. La evidencia estable es el plan y la cantidad de documentos examinados: 3 frente a 12.", color=VERDE)
    doc.add_paragraph(
        "El total observado es 35 índices: ocho _id automáticos en colecciones comunes, un índice meta/fecha creado "
        "automáticamente por la serie temporal y 26 índices definidos por el proyecto. No se debe afirmar que cuantos más "
        "índices mejor: una revisión futura debería eliminar cualquiera que no sostenga una consulta real."
    )

    doc.add_heading("24 GIS y series temporales", level=1)
    doc.add_heading("24 1 Información geográfica", level=2)
    doc.add_paragraph(
        "GIS combina el qué con el dónde. terminales.ubicacion utiliza GeoJSON Point, un modelo vectorial adecuado para "
        "entidades puntuales. MongoDB espera [longitud, latitud], no el orden habitual de lectura latitud–longitud. El índice "
        "2dsphere interpreta la Tierra como superficie esférica y permite $geoNear. Desde un punto cercano a Mendoza, la "
        "consulta devuelve MDZ como terminal más próxima y calcula distancias en metros que el pipeline convierte a kilómetros."
    )
    doc.add_heading("24 2 Serie temporal", level=2)
    doc.add_paragraph(
        "seguimientoviajes usa fecha como timeField y meta como metaField. Meta contiene los identificadores estables de "
        "viaje y colectivo; cada medición agrega ubicación, velocidad, pasajeros estimados y estado. MongoDB agrupa internamente "
        "mediciones en buckets y aplica compresión. El caso cumple los criterios de la clase: inserción ordenada en el tiempo, "
        "consultas por rango y ausencia de correcciones frecuentes sobre datos antiguos."
    )
    add_code(doc, '''db.createCollection("seguimientoviajes", {
  timeseries: {
    timeField: "fecha",
    metaField: "meta",
    granularity: "minutes"
  }
})''')
    add_note(doc, "Por qué reservas no es time series", "Una reserva cambia de pendiente a confirmada, cancelada o vencida. Ese patrón mutable no coincide con una medición append-only; para reservas corresponde una colección común con índice TTL parcial.")

    doc.add_heading("25 Infraestructura, almacenamiento y durabilidad", level=1)
    doc.add_picture(str(ASSETS / "arquitectura.png"), width=Inches(6.8))
    add_caption(doc, "Figura 4. Capas de la demostración local.")
    doc.add_paragraph(
        "mongod es el proceso servidor; mongosh es el cliente de consola; Compass es el cliente gráfico; el servidor Python "
        "expone una API local y delega comandos a mongo_api.js. WiredTiger es el motor de almacenamiento observado. La teoría "
        "de MVCC explica que las lecturas pueden ver versiones consistentes mientras otras operaciones escriben; los checkpoints "
        "crean puntos persistentes y el journal funciona como write-ahead log para durabilidad. Estos mecanismos no reemplazan "
        "un backup: protegen la operación del motor, mientras el respaldo protege ante pérdida lógica o física."
    )
    doc.add_paragraph(
        "Atlas sería la evolución cloud: clúster administrado, replicación, escalado, herramientas de administración y "
        "opciones de backup. La entrega local evita depender de Internet y simplifica la demostración, pero no debe confundirse "
        "con una arquitectura productiva: faltan autenticación del servidor, TLS, secretos, alta disponibilidad y monitoreo."
    )

    doc.add_heading("26 Backup y restauración verificados", level=1)
    doc.add_paragraph(
        "mongoexport produce JSON o CSV para intercambio de datos; no conserva por sí solo todos los tipos BSON, índices, "
        "validadores y opciones. mongodump genera un respaldo BSON apropiado para recuperar la base, y mongorestore lo reconstruye. "
        "El respaldo final se restauró en gestionpasajes_restore_verificacion y se compararon documentos, tipos e índices."
    )
    add_table(doc, ["Colección", "Documentos restaurados", "Índices", "Resultado"], [
        ("colectivos", "5", "3", "coincide"),
        ("pasajeros", "15", "2", "coincide"),
        ("pasajes", "31", "8", "coincide"),
        ("recorridos", "6", "3", "coincide"),
        ("reservas", "9", "5", "coincide"),
        ("seguimientoviajes", "24", "2", "timeseries conservada"),
        ("terminales", "12", "3", "coincide"),
        ("usuarios", "10", "2", "coincide"),
        ("viajes", "12", "7", "coincide"),
    ], widths=[4.2, 4.3, 3.0, 5.0])
    add_code(doc, '''mongodump --db gestionpasajes --out backups/final_2026-10-23
mongorestore --nsFrom='gestionpasajes.*' \\
  --nsTo='gestionpasajes_restore_verificacion.*' \\
  backups/final_2026-10-23''')

    doc.add_heading("27 Guion razonado para la exposición", level=1)
    add_numbered(doc, [
        "Plantear el problema: vender pasajes por tramos sin duplicar la ocupación.",
        "Mostrar viajes como agregado operativo y explicar por qué contiene snapshots e inventario.",
        "Comparar embebido, referencia y snapshot con ejemplos del modelo.",
        "Ejecutar la búsqueda MDZ|COR y explicar el índice compuesto multikey.",
        "Mostrar la disponibilidad por tramo y la regla matemática de superposición.",
        "Explicar la reserva atómica y la compensación si falla la inserción.",
        "Ejecutar un $lookup de manifiesto y una agregación de recaudación.",
        "Mostrar $geoNear y la colección seguimientoviajes con $dateTrunc.",
        "Abrir explain y comparar IXSCAN con COLLSCAN usando documentos examinados.",
        "Cerrar con mongodump/mongorestore y las limitaciones honestas del prototipo.",
    ])
    doc.add_heading("27 1 Preguntas que conviene poder responder", level=2)
    preguntas = [
        ("¿Por qué es NoSQL si hay relaciones?", "Porque el modelo es documental: combina embebido, referencias manuales, arrays y denormalización; las relaciones no desaparecen."),
        ("¿Por qué duplicar datos?", "Para acelerar lecturas y conservar historia; se distingue snapshot inmutable de dato derivado actualizable."),
        ("¿Qué evita una doble reserva?", "El updateOne condicional sobre el documento viaje es atómico y solo hace $push si no existe una ocupación superpuesta."),
        ("¿Qué hace el TTL?", "Elimina en segundo plano reservas pendientes cuyo venceEn pasó; no debe ser la única regla de disponibilidad."),
        ("¿Por qué usar Decimal128?", "Porque el dinero requiere precisión decimal y no errores de representación binaria."),
        ("¿Qué demuestra explain?", "El plan elegido y el trabajo realizado: claves, documentos y tiempo. En esta prueba el índice reduce documentos examinados de 12 a 3."),
        ("¿Por qué la serie temporal es separada?", "La telemetría crece mucho y se consulta por tiempo; separarla evita inflar el documento viaje y permite buckets comprimidos."),
        ("¿Qué no resuelve el prototipo?", "Pagos reales, autenticación productiva, GPS en vivo, alta disponibilidad y monitoreo."),
    ]
    add_table(doc, ["Pregunta", "Respuesta esperada"], preguntas, widths=[5.5, 11.2], font_size=8.2)

    doc.add_heading("28 Conclusión personal", level=1)
    doc.add_paragraph(
        "El aprendizaje principal es que MongoDB no consiste en trasladar tablas a colecciones. El diseño comienza por las "
        "operaciones: qué debe leerse junto, qué necesita atomicidad, qué debe conservar historia y qué crecerá sin límite. "
        "Desde esas preguntas se justifican los arrays, snapshots, referencias, índices, TTL y series temporales. La interfaz "
        "es una prueba del modelo; el núcleo del proyecto sigue siendo la base documental, sus consultas y su operación segura."
    )
    doc.add_heading("28 1 Fuentes de la actualización", level=2)
    add_bullets(doc, [
        "Enunciado Proyecto MongoDB 2026 — Ing. Fernando Zapata.",
        "Presentaciones 0 a 9b de Diseño de Bases de Datos II, Ing. Fernando Zapata, ciclo 2026.",
        "Scripts y evidencia ejecutada de la base gestionpasajes, versión del 23 de octubre de 2026.",
        "Documentación oficial de MongoDB utilizada como ampliación conceptual de las guías de clase.",
    ])

    doc.core_properties.title = "Informe personal de estudio del proyecto MongoDB — gestión de pasajes"
    doc.core_properties.subject = "Diseño, implementación y teoría aplicada de MongoDB"
    doc.core_properties.keywords = "MongoDB, NoSQL, modelado documental, índices, time series, backup"
    OUTPUT_PERSONAL.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PERSONAL)


def portada_profesor(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("UNIVERSIDAD DE MENDOZA\nFacultad de Ingeniería")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(AZUL)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(55)
    r = p.add_run("PROYECTO MONGODB 2026")
    r.bold = True
    r.font.size = Pt(25)
    r.font.color.rgb = RGBColor.from_string(AZUL)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Sistema de gestión de pasajes de larga distancia")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor.from_string(AZUL_MEDIO)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("Modelado e implementación práctica de una solución\nbasada en el motor documental MongoDB")
    r.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(65)
    p.add_run("Materia: Diseño de Bases de Datos II\n").bold = True
    p.add_run("Profesor: Ing. Fernando Zapata\n")
    p.add_run("Estudiante: ______________________________\n")
    p.add_run("Base de datos: gestionpasajes\n")
    p.add_run("Año: 2026")
    doc.add_page_break()


def crear_informe_profesor():
    doc = Document()
    configurar_fuentes(doc)
    configurar_pagina(doc)
    configurar_encabezado_pie(doc, "Proyecto MongoDB 2026 · gestionpasajes")
    portada_profesor(doc)

    doc.add_heading("Resumen ejecutivo", level=1)
    doc.add_paragraph(
        "El proyecto implementa en MongoDB un sistema de gestión de pasajes para una única empresa de colectivos de larga "
        "distancia. Resuelve recorridos con paradas intermedias, venta por tramos, disponibilidad de asientos, reservas "
        "temporales, ventas online y en boletería, pasajeros, pagos simulados, informes y seguimiento temporal. El modelo "
        "prioriza los patrones de consulta de la aplicación y combina documentos embebidos, referencias ObjectId y snapshots "
        "históricos. La entrega incluye nueve colecciones, validadores, datos de prueba, 26 índices definidos por el proyecto, "
        "diez consultas avanzadas, mediciones con explain, una interfaz local y un respaldo BSON cuya restauración fue verificada."
    )
    add_note(doc, "Resultado", "La solución satisface las fases solicitadas en el enunciado 2026: necesidad, modelado, implementación, carga, consultas avanzadas, indización, resguardo y evidencia para la evaluación.", color=VERDE)
    doc.add_heading("Trazabilidad con el enunciado", level=2)
    add_table(doc, ["Requisito", "Implementación", "Evidencia"], [
        ("Necesidad y solución", "gestión integral de pasajes por tramo", "secciones 1 y 2"),
        ("Modelo de datos", "embebido, referencias, snapshots y time series", "secciones 3 y 4"),
        ("Implementación MongoDB", "scripts idempotentes y validadores", "sección 5"),
        ("Carga de datos", "124 documentos lógicos coherentes", "sección 6"),
        ("Consultas avanzadas", "diez consultas ejecutables", "sección 8"),
        ("Desempeño e índices", "explain: IXSCAN frente a COLLSCAN", "sección 9"),
        ("Política de resguardo", "mongodump y restauración verificada", "sección 11"),
    ], widths=[4.5, 7.0, 5.2])

    doc.add_heading("1 Planteo de la necesidad", level=1)
    doc.add_paragraph(
        "Una empresa de transporte necesita centralizar recorridos, horarios, vehículos, disponibilidad y ventas. El problema "
        "distintivo es la venta por tramos: un asiento ocupado Mendoza–San Luis puede venderse nuevamente San Luis–Córdoba, "
        "pero no debe venderse a dos pasajeros en segmentos superpuestos. Además, la compra puede realizarse online o en "
        "boletería, y comprador, usuario y pasajero pueden ser personas diferentes."
    )
    doc.add_heading("1 1 Objetivos", level=2)
    add_bullets(doc, [
        "Buscar viajes por origen, destino y fecha, incluidas paradas intermedias.",
        "Calcular asientos libres por tramo y rechazar ocupaciones superpuestas.",
        "Registrar reservas temporales y ventas confirmadas con pago simulado.",
        "Buscar pasajes por DNI y producir manifiestos e informes de recaudación.",
        "Conservar la fotografía histórica del pasajero, viaje, recorrido y colectivo.",
        "Aplicar validación, índices, consultas geoespaciales, series temporales y respaldo recuperable.",
    ])
    doc.add_heading("1 2 Alcance y exclusiones", level=2)
    add_table(doc, ["Incluido", "Excluido del prototipo"], [
        ("usuarios y roles; pasajeros; terminales; recorridos; colectivos", "integración bancaria real"),
        ("viajes, venta por tramos, reservas y pasajes", "autenticación productiva y publicación en Internet"),
        ("GeoJSON, telemetría simulada e informes", "dispositivos GPS reales y monitoreo continuo"),
        ("interfaz local para demostrar operaciones", "mantenimiento, combustible y encomiendas"),
    ], widths=[8.4, 8.4])

    doc.add_heading("2 Solución y arquitectura", level=1)
    doc.add_picture(str(ASSETS / "arquitectura.png"), width=Inches(6.8))
    add_caption(doc, "Figura 1. Arquitectura local de la solución.")
    doc.add_paragraph(
        "La interfaz web se ejecuta en el navegador. Un servidor HTTP local escrito en Python recibe solicitudes y usa "
        "mongo_api.js mediante mongosh para operar sobre MongoDB en 127.0.0.1:27017. MongoDB 8.0.32 utiliza WiredTiger. "
        "La interfaz no reemplaza la demostración de la base: materializa las consultas y reglas que determinaron el modelo."
    )
    add_table(doc, ["Componente", "Responsabilidad"], [
        ("MongoDB", "persistencia documental, índices, validación y agregaciones"),
        ("mongosh / mongo_api.js", "comandos y operaciones reutilizables"),
        ("server.py", "API HTTP local y serialización JSON"),
        ("interfaz web", "búsqueda, selección de asiento, reservas, pasajeros e informes"),
        ("MongoDB Compass", "inspección gráfica, Explain Plan y My Queries"),
    ], widths=[5.3, 11.5])

    doc.add_heading("3 Criterios de modelado documental", level=1)
    doc.add_paragraph(
        "El modelo se diseñó desde la navegabilidad y las consultas. Se embebe cuando los datos pertenecen al agregado, se "
        "consultan juntos y tienen crecimiento acotado. Se referencia cuando una entidad posee vida independiente o se "
        "comparte. Se crea un snapshot cuando la historia debe permanecer estable aunque cambie el dato maestro."
    )
    add_table(doc, ["Patrón", "Ejemplos", "Motivo"], [
        ("Embebido", "paradas/tramos, asientos, pago, inventario", "lectura conjunta y atomicidad documental"),
        ("Referencia", "viajeId, pasajeroId, colectivoId, recorridoId", "identidad independiente y reutilización"),
        ("Snapshot", "recorridoSnapshot, colectivoSnapshot, pasajeroSnapshot", "historial y consultas sin joins repetidos"),
        ("Documento separado", "reservas y pasajes", "ciclo de vida y consultas propias"),
        ("Time series", "seguimientoviajes", "mediciones append-only ordenadas por fecha"),
    ], widths=[3.4, 6.2, 7.3])
    doc.add_picture(str(ASSETS / "modelo_documental_actualizado.png"), width=Inches(6.8))
    add_caption(doc, "Figura 2. Colecciones, referencias y datos embebidos del modelo implementado.")

    doc.add_page_break()
    doc.add_heading("4 Colecciones y responsabilidades", level=1)
    add_table(doc, ["Colección", "Responsabilidad", "Decisión destacada"], [
        ("usuarios", "cuentas, perfiles y roles", "pasajeroId opcional"),
        ("pasajeros", "personas que viajan", "contacto y domicilio embebidos"),
        ("terminales", "puntos operativos", "GeoJSON Point"),
        ("recorridos", "itinerarios reutilizables", "paradas y tramos embebidos"),
        ("colectivos", "vehículos y distribución", "asientos físicos embebidos"),
        ("viajes", "ejecución de recorrido", "agregado operativo principal"),
        ("reservas", "bloqueos temporales", "TTL parcial para pendientes"),
        ("pasajes", "venta confirmada", "snapshots, pago y embarque"),
        ("seguimientoviajes", "telemetría histórica", "colección time series"),
    ], widths=[3.3, 6.1, 7.4])
    doc.add_heading("4 1 Viajes como agregado operativo", level=2)
    doc.add_paragraph(
        "viajes es la colección principal desde el punto de vista operativo: contiene salida, estado, recorrido y colectivo "
        "fotografiados, pares origen–destino e inventario por asiento. No existe una declaración de colección principal en "
        "MongoDB; su centralidad surge de las operaciones. pasajes es, en cambio, el núcleo transaccional e histórico."
    )
    doc.add_heading("4 2 Regla de ocupación por tramo", level=2)
    doc.add_paragraph(
        "Los tramos se representan con posiciones [desdeOrden, hastaOrden). Dos ocupaciones se superponen cuando el inicio "
        "existente es menor que el fin solicitado y el fin existente es mayor que el inicio solicitado. Al excluir el límite "
        "superior, un asiento puede liberarse en una parada y volver a venderse desde esa misma parada."
    )
    add_code(doc, "existente.desdeOrden < solicitado.hastaOrden\n&& existente.hastaOrden > solicitado.desdeOrden")

    doc.add_heading("5 Implementación física", level=1)
    doc.add_paragraph(
        "scripts/01_estructura.js crea o actualiza la estructura sin borrar datos. Ocho colecciones comunes usan validación "
        "$jsonSchema estricta; la serie temporal se crea con timeField y metaField. Los scripts de carga usan identificadores "
        "fijos y operaciones idempotentes para poder repetirse sin duplicar el conjunto académico."
    )
    add_table(doc, ["Grupo de scripts", "Función"], [
        ("01–05", "estructura, carga base, carga ampliada y verificaciones"),
        ("06–07", "migraciones de DNI y campos optimizados para interfaz"),
        ("08–09", "operaciones de interfaz y prueba de doble reserva"),
        ("10", "creación y carga de la serie temporal"),
        ("11", "catálogo de consultas para presentación"),
        ("12", "verificación integral de colección, referencias, índices, GIS y explain"),
    ], widths=[4.0, 12.7])
    doc.add_heading("5 1 Validación", level=2)
    doc.add_paragraph(
        "Los validadores controlan tipos, campos requeridos, enumeraciones y estructuras anidadas. No sustituyen la integridad "
        "referencial: MongoDB no comprueba por sí solo que un ObjectId exista en otra colección. Esa regla se aplica en las "
        "operaciones y se audita con pipelines de $lookup que detectan referencias huérfanas."
    )
    add_code(doc, '''validator: {
  $jsonSchema: {
    bsonType: "object",
    required: ["codigo", "viajeId", "pasajeroId", "tramo", "estado"],
    properties: { estado: { enum: ["pendiente", "confirmada", "vencida", "cancelada"] } }
  }
},
validationLevel: "strict",
validationAction: "error"''')

    doc.add_heading("6 Carga y estado de los datos", level=1)
    add_table(doc, ["Colección", "Documentos", "Contenido representativo"], [
        ("usuarios", "10", "administradores, vendedores, clientes y conductores"),
        ("pasajeros", "15", "documentos, contacto, correo y domicilio"),
        ("terminales", "12", "Mendoza y destinos nacionales con coordenadas"),
        ("recorridos", "6", "itinerarios directos y con paradas"),
        ("colectivos", "5", "capacidades y mapas de asientos"),
        ("viajes", "12", "salidas programadas e inventarios"),
        ("reservas", "9", "confirmadas y canceladas; pendientes vencidas eliminadas por TTL"),
        ("pasajes", "31", "ventas online/boletería y distintos medios"),
        ("seguimientoviajes", "24", "telemetría simulada de tres viajes"),
    ], widths=[4.0, 2.5, 10.2])
    doc.add_paragraph(
        "Los datos son ficticios y se diseñaron para cubrir ramas de consulta. El TTL parcial ya eliminó reservas pendientes "
        "cuyo vencimiento pasó, mientras conserva confirmadas y canceladas. Esto demuestra que TTL opera en segundo plano y "
        "que el estado debe formar parte de su filtro parcial."
    )

    doc.add_page_break()
    doc.add_heading("7 Operaciones críticas y concurrencia", level=1)
    doc.add_paragraph(
        "La reserva ejecuta updateOne sobre viajes con un filtro que exige ausencia de conflicto y un $push sobre la ocupación "
        "del asiento. La modificación de un documento es atómica. Si otro proceso reservó primero, modifiedCount es cero. "
        "Después se inserta reservas; si esa segunda operación falla, se compensa con $pull. La prueba automatizada crea una "
        "reserva, confirma que el asiento desaparece de disponibles y verifica que un segundo intento sea rechazado."
    )
    add_table(doc, ["Paso", "Operación", "Garantía"], [
        ("1", "leer viaje y posiciones del tramo", "origen anterior a destino"),
        ("2", "updateOne condicional + $push", "bloqueo atómico en el agregado"),
        ("3", "insertOne en reservas", "registro temporal consultable"),
        ("4", "$pull compensatorio si falla insertOne", "evita ocupación huérfana"),
        ("5", "TTL + limpieza lógica", "las reservas vencidas no bloquean disponibilidad"),
    ], widths=[1.5, 7.0, 8.2])

    doc.add_heading("8 Consultas avanzadas", level=1)
    add_table(doc, ["Consulta", "Operadores", "Resultado de prueba"], [
        ("Viajes MDZ–COR", "find, rango, sort", "3 viajes"),
        ("Disponibilidad tramo completo", "$filter, $map, $size", "42 asientos libres en viaje de muestra"),
        ("Pasajes por DNI", "find sobre snapshot", "historial directo"),
        ("Manifiesto", "$lookup, $unwind, $project", "4 pasajeros en viaje de muestra"),
        ("Recaudación", "$match, $group, $sum, $avg", "agrupada por canal y medio"),
        ("Ocupación", "$filter, $size, $divide", "porcentaje por viaje"),
        ("Terminales cercanas", "$geoNear", "MDZ es la más cercana"),
        ("Reservas y TTL", "$group y getIndexes", "estados conservados e índice parcial"),
        ("Velocidad temporal", "$dateTrunc, $avg", "ventanas de 15 minutos"),
        ("Rendimiento", "explain executionStats", "IXSCAN y métricas"),
    ], widths=[5.1, 6.4, 5.3], font_size=8)
    doc.add_heading("8 1 Ejemplo de relación entre colecciones", level=2)
    add_code(doc, '''db.pasajes.aggregate([
  { $match: { viajeId: ObjectId("66f500000000000000000001") } },
  { $lookup: {
      from: "pasajeros",
      localField: "pasajeroId",
      foreignField: "_id",
      as: "pasajeroActual"
  }},
  { $unwind: "$pasajeroActual" },
  { $project: { codigo: 1, asientoNumero: 1, "pasajeroActual.nombre": 1 } }
])''')
    doc.add_paragraph(
        "$lookup relaciona documentos por igualdad de identificadores. Se usa cuando la entidad relacionada necesita leerse "
        "actualizada. Para la búsqueda por DNI no se usa $lookup porque pasajes ya conserva pasajeroSnapshot y posee un índice "
        "sobre el documento; esa es una denormalización intencional."
    )

    doc.add_page_break()
    doc.add_heading("9 Índices y medición de rendimiento", level=1)
    doc.add_paragraph(
        "La base expone 35 índices: ocho _id automáticos, uno meta/fecha automático de la serie temporal y 26 creados por el "
        "proyecto. Incluyen índices únicos, compuestos, multikey, geoespaciales y TTL. Cada uno se vincula con una consulta o "
        "una regla; crear índices indiscriminadamente perjudicaría escrituras y consumo de memoria/disco."
    )
    add_table(doc, ["Índice", "Tipo", "Consulta o regla"], [
        ("uq_pasajeros_documento", "compuesto único", "identidad del pasajero"),
        ("ix_viajes_parVenta_estado_fecha", "compuesto multikey", "origen, destino, estado y fecha"),
        ("ix_pasajes_documento_pasajero", "compuesto", "pasajes por DNI"),
        ("ix_terminales_ubicacion", "2dsphere", "proximidad geográfica"),
        ("ttl_reservas_pendientes_venceEn", "TTL parcial", "eliminar pendientes vencidas"),
        ("ix_seguimiento_viaje_fecha", "compuesto temporal", "mediciones de un viaje por rango"),
    ], widths=[6.8, 4.0, 6.0])
    add_table(doc, ["Plan", "nReturned", "totalKeysExamined", "totalDocsExamined", "Etapa"], [
        ("Con índice", "3", "3", "3", "IXSCAN"),
        ("Forzado sin índice", "3", "0", "12", "COLLSCAN"),
    ], widths=[5.6, 2.5, 3.4, 3.5, 2.0])
    doc.add_paragraph(
        "Con el volumen académico, el tiempo en milisegundos no es una medida estable. La comparación válida es el trabajo: "
        "el índice examina solo los tres documentos que devuelve, mientras el recorrido completo revisa los doce viajes."
    )

    doc.add_heading("10 GIS y series temporales", level=1)
    doc.add_heading("10 1 Terminales geográficas", level=2)
    doc.add_paragraph(
        "terminales.ubicacion almacena puntos GeoJSON con [longitud, latitud]. El índice 2dsphere permite $geoNear, que debe "
        "ser la primera etapa del pipeline. En la prueba desde Mendoza se obtuvieron MDZ, SJN, SRA, LUQ y VME ordenadas por "
        "distancia. La función convierte metros a kilómetros con $divide y redondea con $round."
    )
    doc.add_heading("10 2 Seguimiento temporal", level=2)
    add_code(doc, '''db.createCollection("seguimientoviajes", {
  timeseries: { timeField: "fecha", metaField: "meta", granularity: "minutes" }
})''')
    doc.add_paragraph(
        "La colección contiene 24 mediciones simuladas de tres viajes. meta identifica viaje y colectivo; fecha ordena la "
        "medición; ubicación, velocidad, pasajerosEstimados y estado varían. MongoDB organiza internamente buckets y aplica "
        "compresión. El pipeline de demostración usa $dateTrunc para ventanas de quince minutos y $avg para velocidad y pasajeros."
    )

    doc.add_heading("11 Resguardo y restauración", level=1)
    doc.add_paragraph(
        "Se generó backups/final_2026-10-23 con mongodump. Luego se restauró en la base independiente "
        "gestionpasajes_restore_verificacion. Se compararon conteos, tipos e índices de todas las colecciones. La serie temporal "
        "se recuperó como timeseries y conservó sus 24 mediciones y dos índices. El resultado fue RESTAURACION_VERIFICADA_OK; "
        "la base temporal se eliminó después de la prueba."
    )
    add_code(doc, '''mongodump --db gestionpasajes --out backups/final_2026-10-23
mongorestore --nsFrom='gestionpasajes.*' \\
  --nsTo='gestionpasajes_restore_verificacion.*' \\
  backups/final_2026-10-23''')
    add_table(doc, ["Control", "Resultado"], [
        ("Documentos", "coinciden en las nueve colecciones"),
        ("Índices", "coinciden por colección"),
        ("Validadores", "restaurados desde metadatos"),
        ("Time series", "tipo y mediciones conservados"),
        ("Fallos de restauración", "0"),
    ], widths=[7.0, 9.7])

    doc.add_heading("12 Seguridad, límites y evolución", level=1)
    add_bullets(doc, [
        "Las contraseñas deben almacenarse como hashes generados por la aplicación; el conjunto académico usa valores ficticios.",
        "Las credenciales y cadenas de conexión no se escriben en scripts ni se incluyen en respaldos compartidos.",
        "El servidor local no implementa autenticación HTTP, TLS, rate limiting ni controles de producción.",
        "MongoDB local es adecuado para la defensa; Atlas permitiría clúster, réplicas, escalado, monitoreo y backups administrados.",
        "Si se exigiera atomicidad estricta entre viajes y reservas se usaría una transacción en un replica set.",
        "Las mediciones temporales son simuladas; la integración GPS real requeriría ingestión, retención, seguridad y monitoreo.",
    ])

    doc.add_heading("13 Conclusiones", level=1)
    doc.add_paragraph(
        "La solución demuestra un modelo documental guiado por casos de uso y no una copia de tablas relacionales. viajes "
        "concentra lo necesario para vender y reservar de forma atómica; pasajes conserva el hecho histórico; reservas separa "
        "el ciclo temporal; seguimientoviajes atiende mediciones crecientes; los índices responden a consultas demostrables. "
        "La carga, las operaciones, los pipelines, explain y la restauración fueron ejecutados sobre la instancia local."
    )

    doc.add_heading("Anexo A — Secuencia de demostración", level=1)
    add_numbered(doc, [
        "Abrir gestionpasajes en Compass y mostrar las nueve colecciones.",
        "Abrir viajes y explicar snapshots, paresVenta e inventarioAsientos.",
        "Ejecutar consulta 01 y luego Explain Plan para mostrar IXSCAN.",
        "Ejecutar consulta 02 y explicar disponibilidad por superposición de tramos.",
        "Ejecutar consulta 04 para demostrar $lookup.",
        "Ejecutar consulta 05 para mostrar recaudación con $group.",
        "Ejecutar consulta 07 para mostrar GeoJSON y 2dsphere.",
        "Abrir seguimientoviajes y ejecutar consulta 09 con $dateTrunc.",
        "Mostrar índices de reservas y explicar TTL parcial.",
        "Cerrar con scripts/12_verificacion_completa.js y la evidencia del backup restaurado.",
    ])
    doc.add_heading("Anexo B — Archivos principales", level=1)
    add_table(doc, ["Archivo", "Uso"], [
        ("README.md", "guía de estructura, ejecución y decisiones"),
        ("scripts/01_estructura.js", "colecciones, validadores e índices"),
        ("scripts/02_datos_prueba.js y 04_datos_ampliados.js", "carga reproducible"),
        ("scripts/08_operaciones_interfaz.js", "búsqueda, disponibilidad y reserva atómica"),
        ("scripts/10_series_temporales.js", "telemetría de demostración"),
        ("scripts/11_consultas_presentacion.js", "catálogo de consultas avanzadas"),
        ("scripts/12_verificacion_completa.js", "prueba integral"),
        ("consultas/GUIA_MY_QUERIES.md", "guardado manual en Compass"),
        ("backups/final_2026-10-23", "respaldo BSON restaurable y verificado"),
        ("interfaz", "aplicación local de demostración"),
    ], widths=[7.5, 9.2])

    doc.add_heading("Anexo C — Fuentes académicas", level=1)
    add_bullets(doc, [
        "Proyecto MongoDB 2026 — Implementación práctica de una solución basada en el motor documental MongoDB.",
        "0 — Clases SQL Server.",
        "1 — Introducción a infraestructura 2025.",
        "2 — Clases GIS.",
        "3 — Introducción a bases de datos NoSQL 2025.",
        "4 — Introducción a MongoDB 2025.",
        "5 — MongoDB: instalación y configuración 2025.",
        "6 — MongoDB: modelado 2025.",
        "7 — MongoDB: operación 2025.",
        "8 — MongoDB: índices y rendimiento 2025.",
        "9 — MongoDB: backup 2025.",
        "9b — MongoDB Cloud 2025.",
    ])
    doc.add_paragraph(
        "Las presentaciones funcionan como guía de los temas impartidos. Las explicaciones y decisiones incluidas en este "
        "informe se verificaron mediante la implementación y la ejecución de la base del proyecto."
    )

    doc.core_properties.title = "Informe técnico Proyecto MongoDB 2026 — gestión de pasajes"
    doc.core_properties.subject = "Implementación práctica de MongoDB"
    doc.core_properties.keywords = "MongoDB, pasajes, NoSQL, índices, consultas, backup"
    doc.save(OUTPUT_PROFESOR)


def main():
    asegurar_carpetas()
    crear_diagramas()
    actualizar_informe_personal()
    crear_informe_profesor()
    print(OUTPUT_PERSONAL)
    print(OUTPUT_PROFESOR)


if __name__ == "__main__":
    main()
